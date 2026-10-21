import os
import sys
import subprocess

# Auto-install python-pptx if not present
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("python-pptx is not installed. Installing it now...")
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.shapes import MSO_SHAPE
        print("python-pptx successfully installed!")
    except Exception as e:
        print(f"Failed to install python-pptx: {e}")
        print("Please run: pip install python-pptx")
        sys.exit(1)

def create_presentation():
    prs = Presentation()
    
    # Set to widescreen (16:9)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Design tokens
    COLOR_BG = RGBColor(7, 9, 19)        # Deep space dark #070913
    COLOR_CARD = RGBColor(22, 28, 51)    # Lighter card bg #161c33
    COLOR_TEAL = RGBColor(45, 212, 191)   # #2dd4bf
    COLOR_PURPLE = RGBColor(168, 85, 247) # #a855f7
    COLOR_WHITE = RGBColor(243, 244, 246)  # #f3f4f6
    COLOR_MUTED = RGBColor(156, 163, 175)  # #9ca3af
    
    FONT_TITLE = "Arial"
    FONT_BODY = "Arial"
    
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    def apply_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_header(slide, title_text, tagline_text=None):
        if tagline_text:
            tagline_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
            tf = tagline_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = tagline_text.upper()
            p.font.name = FONT_TITLE
            p.font.size = Pt(10)
            p.font.bold = True
            p.font.color.rgb = COLOR_TEAL
            
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = FONT_TITLE
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE

    def add_card(slide, left, top, width, height, title, body_paragraphs, border_color=None):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD
        if border_color:
            shape.line.color.rgb = border_color
            shape.line.width = Pt(1.5)
        else:
            shape.line.color.rgb = RGBColor(40, 48, 77)
            shape.line.width = Pt(1)
            
        # Text Box inside shape
        text_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.25), width - Inches(0.5), height - Inches(0.5))
        tf = text_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        # Title
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_WHITE
        p_title.space_after = Pt(10)
        
        # Body
        for idx, text in enumerate(body_paragraphs):
            p = tf.add_paragraph() if idx > 0 else tf.paragraphs[1] if len(tf.paragraphs) > 1 else tf.add_paragraph()
            p.text = text
            p.font.name = FONT_BODY
            p.font.size = Pt(11)
            p.font.color.rgb = COLOR_MUTED
            p.space_after = Pt(8)
            
    # ----------------------------------------------------
    # SLIDE 1: COVER
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    apply_background(slide1)
    
    # Giant Logo
    logo_box = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(1.5))
    tf = logo_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "circl"
    p.font.name = FONT_TITLE
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Accent dot in cover logo
    dot_box = slide1.shapes.add_textbox(Inches(7.3), Inches(2.3), Inches(1), Inches(1))
    tf_dot = dot_box.text_frame
    p_dot = tf_dot.paragraphs[0]
    p_dot.text = "."
    p_dot.font.name = FONT_TITLE
    p_dot.font.size = Pt(72)
    p_dot.font.bold = True
    p_dot.font.color.rgb = COLOR_TEAL
    
    # Subtitle
    sub_box = slide1.shapes.add_textbox(Inches(1.5), Inches(3.8), Inches(10.3), Inches(1.0))
    tf_sub = sub_box.text_frame
    tf_sub.word_wrap = True
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "Hyper-Local Resource Sharing & Collaborative Community Tasking"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = COLOR_MUTED
    p_sub.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------
    # SLIDE 2: WHY CIRCL
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    apply_background(slide2)
    add_header(slide2, "Why Circl Was Conceived", "The Challenge & Vision")
    
    # Left Text Area
    left_box = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.5))
    tf_left = left_box.text_frame
    tf_left.word_wrap = True
    tf_left.margin_left = tf_left.margin_top = tf_left.margin_right = tf_left.margin_bottom = 0
    
    p1 = tf_left.paragraphs[0]
    p1.text = "Modern suburban and urban communities suffer from an unintended side effect of development: hyper-isolation and massive material waste."
    p1.font.name = FONT_BODY
    p1.font.size = Pt(14)
    p1.font.color.rgb = COLOR_MUTED
    p1.space_after = Pt(18)
    
    p2 = tf_left.add_paragraph()
    p2.text = "We purchase expensive products that sit idle 99% of the time, and we lack a secure, structured way to interact, exchange value, or ask for simple favours from the people living just 50 meters away."
    p2.font.name = FONT_BODY
    p2.font.size = Pt(14)
    p2.font.color.rgb = COLOR_MUTED
    p2.space_after = Pt(18)

    # Right Cards
    add_card(
        slide2, 
        left=Inches(6.8), 
        top=Inches(1.8), 
        width=Inches(5.7), 
        height=Inches(2.2), 
        title="Material Underutilization",
        body_paragraphs=[
            "An average domestic power drill is used for only 13 minutes in its entire lifespan.",
            "Circl unlocks these dormant community inventories to decrease waste."
        ],
        border_color=COLOR_PURPLE
    )
    
    add_card(
        slide2, 
        left=Inches(6.8), 
        top=Inches(4.3), 
        width=Inches(5.7), 
        height=Inches(2.2), 
        title="The \"Stranger Danger\" Friction",
        body_paragraphs=[
            "People want to help and share, but fear damage, theft, or scams.",
            "Circl builds a dedicated trust engine to facilitate hyper-local sharing safely."
        ],
        border_color=COLOR_TEAL
    )

    # ----------------------------------------------------
    # SLIDE 3: CORE OFFERING
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    apply_background(slide3)
    add_header(slide3, "A Dual-Pillared Marketplace", "The Product Ecosystem")
    
    # Card 1: Resource Sharing
    add_card(
        slide3,
        left=Inches(0.8),
        top=Inches(1.8),
        width=Inches(5.6),
        height=Inches(4.7),
        title="Physical Resource Sharing",
        body_paragraphs=[
            "• Borrowing (Goodwill): Free, neighborly lending of basic household items, books, and appliances to build local rapport.",
            "• Renting (Monetized): Secure hire of high-value tools, vehicles, or gear with system-managed financial deposit protections."
        ],
        border_color=COLOR_TEAL
    )
    
    # Card 2: Local Services
    add_card(
        slide3,
        left=Inches(6.9),
        top=Inches(1.8),
        width=Inches(5.6),
        height=Inches(4.7),
        title="Collaborative Local Services",
        body_paragraphs=[
            "• Friendly Favours: Volunteer aid (pet-feeding, plant-watering, elderly check-ins) to support the neighborhood network.",
            "• Micro-tasks & Gigs: Paid assignments (yard cleaning, technical setups, local delivery) creating gig economy jobs right at home."
        ],
        border_color=COLOR_PURPLE
    )

    # ----------------------------------------------------
    # SLIDE 4: TECH STACK
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    apply_background(slide4)
    add_header(slide4, "The Technological Toolset", "Under the Hood")
    
    sub_stack_box = slide4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.5))
    p_stack = sub_stack_box.text_frame.paragraphs[0]
    p_stack.text = "Circl is built on a highly modern, reactive stack optimized for lightning-fast mobile performance, geolocation, and security."
    p_stack.font.name = FONT_BODY
    p_stack.font.size = Pt(13)
    p_stack.font.color.rgb = COLOR_MUTED

    card_w = Inches(2.2)
    card_h = Inches(4.3)
    card_y = Inches(2.2)
    gap = Inches(0.18)
    start_x = Inches(0.8)
    
    stack_items = [
        {"title": "Flutter", "desc": "Cross-platform client engine delivering smooth, high-fidelity UI animations on iOS and Android."},
        {"title": "Supabase", "desc": "PostgreSQL database, real-time chat sync, file storage, and secure authentication backend."},
        {"title": "Google Maps", "desc": "Geospatial matching, live distance calculations, and interactive map interfaces."},
        {"title": "PayHere SDK", "desc": "Mobile SDK handling secure checkout, credit cards, and escrow safety holds."},
        {"title": "ML Kit OCR", "desc": "Extracts text and validates authenticity of national identity cards (NICs) instantly."}
    ]
    
    for i, item in enumerate(stack_items):
        add_card(
            slide4,
            left=start_x + i * (card_w + gap),
            top=card_y,
            width=card_w,
            height=card_h,
            title=item["title"],
            body_paragraphs=[item["desc"]]
        )

    # ----------------------------------------------------
    # SLIDE 5: TRUST FRAMEWORK
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    apply_background(slide5)
    add_header(slide5, "Aura: The Community Reputation Engine", "The Trust Framework")
    
    sub_trust_box = slide5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.5))
    p_trust = sub_trust_box.text_frame.paragraphs[0]
    p_trust.text = "To eliminate friction, Circl features a complete validation, safety escrow, and gamification framework."
    p_trust.font.name = FONT_BODY
    p_trust.font.size = Pt(13)
    p_trust.font.color.rgb = COLOR_MUTED

    step_w = Inches(3.7)
    step_h = Inches(4.3)
    step_y = Inches(2.2)
    step_gap = Inches(0.3)
    step_start_x = Inches(0.8)
    
    steps = [
        {
            "num": "01",
            "title": "NIC Verification",
            "desc": "Google ML Kit OCR processes government-issued cards, ensuring every profile maps to a verified real-world individual."
        },
        {
            "num": "02",
            "title": "Secure Escrow",
            "desc": "High-value tool rentals hold a pre-authorized security deposit. Money is released only after successful return and inspection."
        },
        {
            "num": "03",
            "title": "Aura Point Economy",
            "desc": "Users earn points and level up their 'Aura' by completing tasks, sharing items, and accumulating high ratings. High Aura unlocks reduced deposit requirements and premium badges."
        }
    ]
    
    for i, step in enumerate(steps):
        add_card(
            slide5,
            left=step_start_x + i * (step_w + step_gap),
            top=step_y,
            width=step_w,
            height=step_h,
            title=f"{step['num']}  {step['title']}",
            body_paragraphs=[step["desc"]],
            border_color=COLOR_TEAL if i % 2 == 0 else COLOR_PURPLE
        )

    # ----------------------------------------------------
    # SLIDE 6: CONCLUSION
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    apply_background(slide6)
    add_header(slide6, "The Local Future is Circular", "Looking Forward")
    
    # Left Text Area
    left_box_6 = slide6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.5))
    tf_left_6 = left_box_6.text_frame
    tf_left_6.word_wrap = True
    tf_left_6.margin_left = tf_left_6.margin_top = tf_left_6.margin_right = tf_left_6.margin_bottom = 0
    
    p6_1 = tf_left_6.paragraphs[0]
    p6_1.text = "By turning physical neighborhoods into active, collaborative sharing networks, Circl decreases carbon footprints, saves users money, and fosters resilient social bonds."
    p6_1.font.name = FONT_BODY
    p6_1.font.size = Pt(16)
    p6_1.font.color.rgb = COLOR_WHITE
    p6_1.space_after = Pt(20)
    
    p6_2 = tf_left_6.add_paragraph()
    p6_2.text = "It's not just an app—it's a paradigm shift in how we own, work, and interact right outside our front door."
    p6_2.font.name = FONT_BODY
    p6_2.font.size = Pt(14)
    p6_2.font.color.rgb = COLOR_MUTED
    p6_2.space_after = Pt(20)

    # Right Metrics Grid (4 items)
    metrics = [
        {"val": "95%", "lbl": "Reduction in idle household item waste"},
        {"val": "$0", "lbl": "Cost to perform goodwill favors and tasks"},
        {"val": "100%", "lbl": "Verified profiles for worry-free interactions"},
        {"val": "15m", "lbl": "Average distance to matching resources"}
    ]
    
    mw = Inches(2.7)
    mh = Inches(2.1)
    mx = Inches(6.8)
    my = Inches(1.8)
    mgap_x = Inches(0.3)
    mgap_y = Inches(0.3)
    
    for idx, metric in enumerate(metrics):
        col = idx % 2
        row = idx // 2
        
        card_x = mx + col * (mw + mgap_x)
        card_y = my + row * (mh + mgap_y)
        
        # Draw metric card background
        shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_x, card_y, mw, mh)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD
        shape.line.color.rgb = RGBColor(40, 48, 77)
        shape.line.width = Pt(1)
        
        # Add metric value
        val_box = slide6.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(0.2), mw - Inches(0.4), Inches(0.8))
        tf_val = val_box.text_frame
        tf_val.word_wrap = True
        tf_val.margin_left = tf_val.margin_top = tf_val.margin_right = tf_val.margin_bottom = 0
        p_val = tf_val.paragraphs[0]
        p_val.text = metric["val"]
        p_val.font.name = FONT_TITLE
        p_val.font.size = Pt(36)
        p_val.font.bold = True
        p_val.font.color.rgb = COLOR_TEAL
        
        # Add metric label
        lbl_box = slide6.shapes.add_textbox(card_x + Inches(0.2), card_y + Inches(1.0), mw - Inches(0.4), Inches(0.9))
        tf_lbl = lbl_box.text_frame
        tf_lbl.word_wrap = True
        tf_lbl.margin_left = tf_lbl.margin_top = tf_lbl.margin_right = tf_lbl.margin_bottom = 0
        p_lbl = tf_lbl.paragraphs[0]
        p_lbl.text = metric["lbl"]
        p_lbl.font.name = FONT_BODY
        p_lbl.font.size = Pt(11)
        p_lbl.font.color.rgb = COLOR_MUTED

    # Save presentation
    output_filename = "circl_presentation.pptx"
    prs.save(output_filename)
    print(f"Presentation saved successfully as '{output_filename}'!")

if __name__ == "__main__":
    create_presentation()
